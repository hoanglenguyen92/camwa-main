import io
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional

from fastapi import BackgroundTasks, HTTPException, status
from openpyxl.reader.excel import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook import Workbook
from sqlalchemy import ColumnElement, select
from sqlalchemy.ext.asyncio import AsyncSession

from backend.core.keycloak import keycloak
from backend.helpers.http import reject_rows_outside_program
from backend.helpers.pagination import Page, PageParams, paginate
from backend.models import Attendance, AttendanceRequest, Exam, Module, ModuleRegistration, Student
from backend.models.enums import AttendanceStatus, RequestStatus
from backend.schemas.attendance import (
    AttendanceCreate,
    AttendanceRequestCreate,
    AttendanceRequestResponse,
    AttendanceResponse,
    AttendanceUpdate,
    CorrectionDecision,
    ExamEligibilityResponse,
)
from backend.services.email_service import send_attendance_correction_notification, send_attendance_request_confirmation
from backend.services.notification_service import create_new_request_notification, create_request_processed_notification

logger = logging.getLogger(__name__)

PRESENT_STATUSES = {AttendanceStatus.PRESENT, AttendanceStatus.LATE, AttendanceStatus.EXCUSED}
ELIGIBILITY_THRESHOLD = 80.0


'''Attendance'''

async def submit_attendance(session: AsyncSession, data: AttendanceCreate) -> AttendanceResponse:
    reg_stmt = select(ModuleRegistration).where(
        ModuleRegistration.student_id == data.student_id,
        ModuleRegistration.module_id == data.module_id,
    )
    if not (await session.execute(reg_stmt)).scalars().first():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Student {data.student_id} is not registered for module {data.module_id}",
        )

    new_att = Attendance(**data.model_dump())
    session.add(new_att)
    await session.commit()
    await session.refresh(new_att)
    return AttendanceResponse.model_validate(new_att)


async def list_attendance(
    session: AsyncSession,
    page: PageParams,
    scope: ColumnElement[bool],
    student_id: Optional[str] = None,
    module_id: Optional[str] = None,
) -> Page[AttendanceResponse]:
    stmt = select(Attendance).where(scope).order_by(Attendance.created_at.desc(), Attendance.attendance_id)
    if student_id:
        stmt = stmt.where(Attendance.student_id == student_id)
    if module_id:
        stmt = stmt.where(Attendance.module_id == module_id)
    return await paginate(session, stmt, AttendanceResponse, page)


async def update_attendance(session: AsyncSession, attendance: Attendance, data: AttendanceUpdate) -> AttendanceResponse:
    attendance.attendance_status = data.attendance_status
    await session.commit()
    await session.refresh(attendance)
    return AttendanceResponse.model_validate(attendance)


async def delete_attendance(session: AsyncSession, attendance: Attendance) -> None:
    await session.delete(attendance)
    await session.commit()


async def import_attendance_from_excel(
    session: AsyncSession, file_bytes: bytes, allowed_program_id: Optional[str] = None
):
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    successful = []
    failed = []
    outside_program = []
    program_of_module: Dict[str, Optional[str]] = {}
    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        student_id = str(row[0]).strip() if len(row) > 0 and row[0] else None
        module_id = str(row[1]).strip() if len(row) > 1 and row[1] else None
        status_value = str(row[3]).strip() if len(row) > 3 and row[3] else None
        if not student_id or not module_id or not status_value:
            failed.append({"row": row_index, "error": "Missing required fields"})
            continue
        if allowed_program_id:
            if module_id not in program_of_module:
                program_of_module[module_id] = await session.scalar(
                    select(Module.program_id).where(Module.module_id == module_id)
                )
            if program_of_module[module_id] != allowed_program_id:
                outside_program.append(row_index)
                continue
        try:
            attendance_status = AttendanceStatus(status_value.upper())
        except ValueError:
            failed.append({
                "row": row_index,
                "student_id": student_id,
                "module_id": module_id,
                "error": f"Invalid attendance status: {status_value}. Must be PRESENT, ABSENT, LATE, or EXCUSED",
            })
            continue
        reg_stmt = select(ModuleRegistration).where(
            ModuleRegistration.student_id == student_id,
            ModuleRegistration.module_id == module_id,
        )
        if not (await session.execute(reg_stmt)).scalars().first():
            failed.append({
                "row": row_index,
                "student_id": student_id,
                "module_id": module_id,
                "error": "Student not registered",
            })
            continue
        new_attendance = Attendance(student_id=student_id, module_id=module_id, attendance_status=attendance_status)
        session.add(new_attendance)
        await session.flush()
        successful.append({
            "attendance_id": new_attendance.attendance_id,
            "student_id": student_id,
            "module_id": module_id,
            "attendance_status": attendance_status.value,
        })
    if outside_program:
        await session.rollback()
        reject_rows_outside_program(outside_program)
    await session.commit()
    return {"successful": successful, "failed": failed}


'''Attendance requests'''

async def list_requests(
    session: AsyncSession,
    page: PageParams,
    scope: ColumnElement[bool],
    request_status: Optional[RequestStatus] = None,
    student_id: Optional[str] = None,
    module_id: Optional[str] = None,
) -> Page[AttendanceRequestResponse]:
    stmt = select(AttendanceRequest).where(scope).order_by(AttendanceRequest.created_at.desc())
    if request_status:
        stmt = stmt.where(AttendanceRequest.request_status == request_status)
    if student_id:
        stmt = stmt.where(AttendanceRequest.student_id == student_id)
    if module_id:
        stmt = stmt.where(AttendanceRequest.module_id == module_id)
    return await paginate(session, stmt, AttendanceRequestResponse, page)


async def delete_attendance_request(session: AsyncSession, req: AttendanceRequest) -> None:
    await session.delete(req)
    await session.commit()


async def _get_student_email(student_id: str) -> Optional[str]:
    try:
        account = await keycloak.get_user(student_id)
    except HTTPException:
        logger.exception("Keycloak lookup failed for student %s", student_id)
        return None
    return account.email if account else None


async def request_correction(
    session: AsyncSession,
    background_tasks: BackgroundTasks,
    data: AttendanceRequestCreate,
) -> AttendanceRequestResponse:
    attendance = await session.get(Attendance, data.attendance_id)
    if not attendance or attendance.student_id != data.student_id or attendance.module_id != data.module_id:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Attendance record not found")

    pending_stmt = select(AttendanceRequest).where(
        AttendanceRequest.attendance_id == data.attendance_id,
        AttendanceRequest.student_id == data.student_id,
        AttendanceRequest.request_status == RequestStatus.PENDING,
    )
    if (await session.execute(pending_stmt)).scalars().first():
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="There is already a pending correction request for this attendance record",
        )

    new_req = AttendanceRequest(**data.model_dump(), request_status=RequestStatus.PENDING)
    session.add(new_req)
    await session.commit()
    await session.refresh(new_req)

    student_email = await _get_student_email(data.student_id)
    if student_email:
        background_tasks.add_task(
            send_attendance_request_confirmation,
            student_email=student_email,
            request_data={
                "module_id": data.module_id,
                "attendance_date": attendance.created_at,
                "current_status": attendance.attendance_status,
                "proposed_status": data.proposed_status,
                "reason": data.reason,
            },
            request_id=new_req.request_id,
        )
    else:
        logger.warning("Could not find email for student %s - confirmation email not sent", data.student_id)

    background_tasks.add_task(create_new_request_notification, session, new_req.request_id, data.student_id)
    return AttendanceRequestResponse.model_validate(new_req)


async def decide_correction(
    session: AsyncSession,
    background_tasks: BackgroundTasks,
    req: AttendanceRequest,
    decision: CorrectionDecision,
    processed_by: str,
) -> AttendanceRequestResponse:
    if req.request_status != RequestStatus.PENDING:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="This request has already been processed")

    is_approved = decision.status == RequestStatus.APPROVED
    approved_status = (decision.approved_status or req.proposed_status) if is_approved else None

    original_attendance = await session.get(Attendance, req.attendance_id)
    original_status = original_attendance.attendance_status if original_attendance else None

    now = datetime.now()
    req.request_status = decision.status
    req.approved_status = approved_status
    req.processed_by = processed_by
    req.processed_at = now
    if is_approved and original_attendance:
        original_attendance.attendance_status = approved_status
        original_attendance.updated_at = now
    logger.info("Correction request %s %s by %s", req.request_id, decision.status.value, processed_by)

    await session.commit()
    await session.refresh(req)

    student_email = await _get_student_email(req.student_id)
    if student_email:
        background_tasks.add_task(
            send_attendance_correction_notification,
            student_email=student_email,
            request_data={
                "module_id": req.module_id,
                "original_status": original_status,
                "proposed_status": req.proposed_status,
                "approved_status": approved_status,
                "reason": req.reason,
            },
            is_approved=is_approved,
            processed_by=processed_by,
        )
    else:
        logger.warning("Could not find email for student %s - correction email not sent", req.student_id)

    background_tasks.add_task(
        create_request_processed_notification, session, req.request_id, processed_by, is_approved
    )
    return AttendanceRequestResponse.model_validate(req)


'''Attendance rate and exam eligibility'''

async def calculate_attendance_rate(session: AsyncSession, student_id: str, module_id: str) -> float:
    stmt = select(Attendance.attendance_status).where(
        Attendance.student_id == student_id,
        Attendance.module_id == module_id,
    )
    statuses = (await session.execute(stmt)).scalars().all()
    if not statuses:
        return 0.0
    present = sum(1 for s in statuses if s in PRESENT_STATUSES)
    return round(present / len(statuses) * 100, 2)


async def check_exam_eligibility(session: AsyncSession, student_id: str, module_id: str) -> ExamEligibilityResponse:
    reg_stmt = select(ModuleRegistration).where(
        ModuleRegistration.student_id == student_id,
        ModuleRegistration.module_id == module_id,
    )
    if not (await session.execute(reg_stmt)).scalars().first():
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Student {student_id} is not registered for module {module_id}",
        )
    rate = await calculate_attendance_rate(session, student_id, module_id)
    return ExamEligibilityResponse(
        student_id=student_id,
        module_id=module_id,
        attendance_rate=rate,
        is_eligible=rate >= ELIGIBILITY_THRESHOLD,
    )


async def _registrations_in_reach(session: AsyncSession, module_ids) -> List[ModuleRegistration]:
    stmt = (
        select(ModuleRegistration)
        .where(ModuleRegistration.module_id.in_(module_ids))
        .order_by(ModuleRegistration.module_id, ModuleRegistration.student_id)
    )
    return list((await session.execute(stmt)).scalars().all())


async def recalculate_exam_eligibility(session: AsyncSession, module_ids) -> Dict[str, Any]:
    """Writes one Exam row per registration in the given modules and summarises per module."""
    registrations = await _registrations_in_reach(session, module_ids)
    existing = {
        (e.student_id, e.module_id): e
        for e in (await session.execute(select(Exam).where(Exam.module_id.in_(module_ids)))).scalars().all()
    }

    modules: Dict[str, Dict[str, Any]] = {}
    for reg in registrations:
        rate = await calculate_attendance_rate(session, reg.student_id, reg.module_id)
        is_eligible = rate >= ELIGIBILITY_THRESHOLD
        exam = existing.get((reg.student_id, reg.module_id))
        if exam:
            exam.attendance_rate = rate
            exam.is_eligible = is_eligible
        else:
            session.add(Exam(student_id=reg.student_id, module_id=reg.module_id, attendance_rate=rate, is_eligible=is_eligible))
        summary = modules.setdefault(reg.module_id, {"module_id": reg.module_id, "eligible": 0, "ineligible": 0})
        summary["eligible" if is_eligible else "ineligible"] += 1
    await session.commit()

    eligible = sum(m["eligible"] for m in modules.values())
    return {
        "modules_processed": len(modules),
        "students_eligible": eligible,
        "students_ineligible": len(registrations) - eligible,
        "modules": list(modules.values()),
    }


async def get_exam_eligibility_by_module(session: AsyncSession, module_id: str) -> List[Exam]:
    stmt = select(Exam).where(Exam.module_id == module_id).order_by(Exam.student_id)
    return list((await session.execute(stmt)).scalars().all())


async def get_exam_eligibility_by_student(session: AsyncSession, student_id: str) -> List[Exam]:
    stmt = select(Exam).where(Exam.student_id == student_id).order_by(Exam.module_id)
    return list((await session.execute(stmt)).scalars().all())


'''Excel export'''

_HEADER_FILL = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
_ELIGIBLE_FILL = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
_INELIGIBLE_FILL = PatternFill(start_color="FFCCCB", end_color="FFCCCB", fill_type="solid")
_EXPORT_COLUMNS = [
    ("Module ID", 15),
    ("Module Name", 30),
    ("Lecturer ID", 15),
    ("Student ID", 15),
    ("Student Name", 25),
    ("Attendance Rate", 18),
    ("Is Eligible", 12),
]


async def export_exam_eligibility(session: AsyncSession, module_ids) -> bytes:
    """One sheet per module, rates computed from current attendance. Nothing is written."""
    stmt = (
        select(Module, Student)
        .join(ModuleRegistration, ModuleRegistration.module_id == Module.module_id)
        .join(Student, Student.student_id == ModuleRegistration.student_id)
        .where(Module.module_id.in_(module_ids))
        .order_by(Module.module_id, Student.student_id)
    )
    rows = (await session.execute(stmt)).all()

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets: Dict[str, Any] = {}
    for module, student in rows:
        sheet = sheets.get(module.module_id)
        if sheet is None:
            sheet = workbook.create_sheet(title=f"{module.module_id}"[:31])
            sheet.append([name for name, _ in _EXPORT_COLUMNS])
            for cell, (_, width) in zip(sheet[1], _EXPORT_COLUMNS):
                cell.font = Font(bold=True)
                cell.fill = _HEADER_FILL
                sheet.column_dimensions[cell.column_letter].width = width
            sheets[module.module_id] = sheet
        rate = await calculate_attendance_rate(session, student.student_id, module.module_id)
        eligible = rate >= ELIGIBILITY_THRESHOLD
        sheet.append([
            module.module_id,
            module.name,
            module.lecturer_id,
            student.student_id,
            student.name,
            f"{rate:.2f}%",
            "YES" if eligible else "NO",
        ])
        sheet.cell(row=sheet.max_row, column=len(_EXPORT_COLUMNS)).fill = _ELIGIBLE_FILL if eligible else _INELIGIBLE_FILL
    if not sheets:
        workbook.create_sheet(title="Exam eligibility").append([name for name, _ in _EXPORT_COLUMNS])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
